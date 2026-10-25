"""Import prospects from Excel (.xlsx) or CSV.

Flow (stateless): the client uploads a file to `parse_upload`, gets the header
row and raw rows back, picks a column mapping, then sends the mapped rows to
`validate_rows` (dry run) and `import_rows` (all-or-nothing transaction).
Each row is validated with the same serializer as the Add Prospect form.
"""

import csv
import io
import re
from datetime import date, datetime

from django.db import transaction
from rest_framework.exceptions import ValidationError

from apps.audit.models import AuditAction
from apps.audit.services import record
from apps.common.normalize import normalize_email, normalize_phone

from .models import ProductCategory, Prospect, ProspectSource
from .statuses import ProspectStatus

MAX_ROWS = 2000
MAX_COLUMNS = 60
MAX_BYTES = 4 * 1024 * 1024

# Target fields the user can map columns to, with header aliases used for guessing.
FIELDS: dict[str, list[str]] = {
    "full_name": ["name", "full name", "fullname", "prospect", "prospect name", "client name", "pangalan"],
    "first_name": ["first name", "firstname", "given name", "first"],
    "last_name": ["last name", "lastname", "surname", "family name", "last"],
    "preferred_name": ["preferred name", "nickname", "nick name", "alias"],
    "mobile": ["mobile", "mobile number", "phone", "phone number", "contact", "contact number", "cellphone", "cell", "cp", "number", "tel"],
    "email": ["email", "e-mail", "email address", "mail"],
    "birth_date": ["birth date", "birthdate", "birthday", "date of birth", "dob"],
    "occupation": ["occupation", "job", "job title", "position", "profession", "work"],
    "company": ["company", "employer", "organization", "office"],
    "location": ["location", "city", "address", "area", "province"],
    "source": ["source", "lead source", "prospect source", "channel"],
    "referral_name": ["referral", "referred by", "referrer", "referral name"],
    "insurance_interests": ["interest", "interests", "insurance interest", "product", "products", "needs"],
    "existing_coverage": ["existing coverage", "existing insurance", "current insurance", "coverage"],
    "priority": ["priority"],
    "status": ["status", "stage", "pipeline stage"],
    "notes": ["notes", "note", "remarks", "comments"],
    "consent_obtained": ["consent", "consent obtained", "agreed to contact"],
}

STATUS_ALIASES = {
    "new": "new_lead", "lead": "new_lead", "attempted": "contact_attempted", "no answer": "contact_attempted",
    "presentation": "presentation_scheduled", "follow-up": "nurturing", "follow up": "nurturing",
    "nurture": "nurturing", "application": "application_in_progress", "submitted": "application_submitted",
    "won": "won", "issued": "won", "policy issued": "won", "client": "won", "lost": "lost",
    "not interested": "lost", "hold": "on_hold",
}


def _norm_header(value) -> str:
    return re.sub(r"[\s_\-]+", " ", str(value or "")).strip().lower()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores phone numbers as numbers
    return str(value).strip()


def parse_upload(upload) -> tuple[list[str], list[list[str]]]:
    """Return (headers, rows) from an uploaded .xlsx or .csv file."""
    name = (upload.name or "").lower()
    if upload.size > MAX_BYTES:
        raise ValidationError({"file": ["The file is larger than 4 MB. Split it into smaller files."]})
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError({"file": ["This Excel file couldn't be read. Save it as .xlsx or CSV and try again."]}) from exc
        sheet = workbook.worksheets[0]
        table = [[_cell(c) for c in row[:MAX_COLUMNS]] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith(".csv") or name.endswith(".txt"):
        raw = upload.read()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        table = [[c.strip() for c in row[:MAX_COLUMNS]] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValidationError({"file": ["Upload an Excel (.xlsx) or CSV file. Older .xls files: save as .xlsx first."]})

    table = [row for row in table if any(cell for cell in row)]
    if not table:
        raise ValidationError({"file": ["The file is empty."]})
    headers = [h or f"Column {i + 1}" for i, h in enumerate(table[0])]
    rows = [row + [""] * (len(headers) - len(row)) for row in table[1:]]
    if len(rows) > MAX_ROWS:
        raise ValidationError({"file": [f"The file has {len(rows)} rows; the limit is {MAX_ROWS} per import."]})
    return headers, rows


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Map column index (as string) to a target field, based on header names."""
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for index, header in enumerate(headers):
        key = _norm_header(header)
        for field, aliases in FIELDS.items():
            if field not in used and (key == field.replace("_", " ") or key in aliases):
                mapping[str(index)] = field
                used.add(field)
                break
    # A full-name column only makes sense when there are no separate first/last columns.
    if "full_name" in used and ({"first_name", "last_name"} & used):
        mapping = {k: v for k, v in mapping.items() if v != "full_name"}
    return mapping


# Surname particles common in Filipino and Spanish names ("Juan Dela Cruz", "Ana De Los Reyes").
SURNAME_PARTICLES = {"de", "dela", "del", "delos", "los", "las", "la", "da", "di", "san", "sta", "sta.", "santa", "van", "von"}


def split_full_name(full: str) -> tuple[str, str]:
    parts = full.split()
    if len(parts) < 2:
        return (parts[0] if parts else ""), ""
    last = [parts[-1]]
    i = len(parts) - 2
    while i > 0 and parts[i].lower() in SURNAME_PARTICLES:
        last.insert(0, parts[i])
        i -= 1
    return " ".join(parts[: i + 1]), " ".join(last)


def _parse_date(value: str) -> str | None:
    value = value.strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    return "invalid"


def _status(value: str) -> str | None:
    key = value.strip().lower()
    if not key:
        return None
    for choice in ProspectStatus:
        if key in {choice.value, choice.label.lower()}:
            return choice.value
    return STATUS_ALIASES.get(key, "invalid")


class _Lookups:
    def __init__(self, user):
        self.user = user
        self.sources = {s.name.lower(): s for s in ProspectSource.objects.filter(owner=user)}
        self.categories = {c.name.lower(): c for c in ProductCategory.objects.filter(owner=user)}
        self.new_sources: set[str] = set()


def map_record(raw: dict[str, str], lookups: _Lookups) -> tuple[dict, list[str]]:
    """Turn one mapped row into serializer input. Returns (data, warnings)."""
    warnings: list[str] = []
    data: dict = {}
    first, last = raw.get("first_name", "").strip(), raw.get("last_name", "").strip()
    if not (first or last) and raw.get("full_name"):
        first, last = split_full_name(raw["full_name"])
    data["first_name"], data["last_name"] = first, last
    for field in ["preferred_name", "mobile", "email", "occupation", "company", "location", "referral_name",
                  "existing_coverage", "notes"]:
        if raw.get(field):
            data[field] = raw[field].strip()

    if raw.get("birth_date"):
        parsed = _parse_date(raw["birth_date"])
        data["birth_date"] = parsed if parsed != "invalid" else raw["birth_date"]
    if raw.get("priority"):
        priority = raw["priority"].strip().lower()
        data["priority"] = {"h": "high", "m": "medium", "l": "low", "urgent": "high"}.get(priority, priority)
    if raw.get("status"):
        status = _status(raw["status"])
        if status == "invalid":
            warnings.append(f"Unknown status “{raw['status']}” — imported as New Lead.")
        else:
            data["status"] = status
    if raw.get("consent_obtained"):
        data["consent_obtained"] = raw["consent_obtained"].strip().lower() in {"yes", "y", "true", "1", "oo"}

    if raw.get("source"):
        name = raw["source"].strip()
        source = lookups.sources.get(name.lower())
        if source:
            data["source"] = source.pk
        else:
            data["_new_source"] = name[:80]
            warnings.append(f"New source “{name}” will be added to your list.")
    if raw.get("insurance_interests"):
        ids, unknown = [], []
        for part in re.split(r"[,;/]", raw["insurance_interests"]):
            part = part.strip()
            if not part:
                continue
            category = lookups.categories.get(part.lower())
            (ids.append(category.pk) if category else unknown.append(part))
        data["insurance_interests"] = ids
        if unknown:
            warnings.append(f"Unrecognised interests skipped: {', '.join(unknown)}.")
    return data, warnings


def _check(records: list[dict], mapping: dict[str, str], request) -> list[dict]:
    from .serializers import ProspectWriteSerializer

    lookups = _Lookups(request.user)
    existing_phones = dict(
        Prospect.objects.filter(owner=request.user).exclude(mobile_normalized="").values_list("mobile_normalized", "id")
    )
    existing_emails = dict(
        Prospect.objects.filter(owner=request.user).exclude(email_normalized="").values_list("email_normalized", "id")
    )
    seen_phones: dict[str, int] = {}
    seen_emails: dict[str, int] = {}
    results = []
    for number, row in enumerate(records, start=2):  # row 1 is the header
        cells = row if isinstance(row, list) else []
        raw = {field: str(cells[int(index)]) if int(index) < len(cells) and cells[int(index)] is not None else ""
               for index, field in mapping.items() if field in FIELDS}
        data, warnings = map_record(raw, lookups)
        new_source = data.pop("_new_source", None)
        serializer = ProspectWriteSerializer(data=data, context={"request": request})
        valid = serializer.is_valid()
        errors = {k: [str(m) for m in v] for k, v in serializer.errors.items()} if not valid else {}

        duplicate_of = None
        phone, email = normalize_phone(data.get("mobile")), normalize_email(data.get("email"))
        if phone and phone in existing_phones:
            duplicate_of = {"type": "existing", "prospect_id": existing_phones[phone], "matched_on": "mobile"}
        elif email and email in existing_emails:
            duplicate_of = {"type": "existing", "prospect_id": existing_emails[email], "matched_on": "email"}
        elif phone and phone in seen_phones:
            duplicate_of = {"type": "file", "row": seen_phones[phone], "matched_on": "mobile"}
        elif email and email in seen_emails:
            duplicate_of = {"type": "file", "row": seen_emails[email], "matched_on": "email"}
        if phone:
            seen_phones.setdefault(phone, number)
        if email:
            seen_emails.setdefault(email, number)

        results.append({
            "row": number,
            "name": f"{data.get('first_name', '')} {data.get('last_name', '')}".strip(),
            "valid": valid,
            "errors": errors,
            "warnings": warnings,
            "duplicate_of": duplicate_of,
            "_data": data,
            "_new_source": new_source,
        })
    return results


def validate_rows(rows: list, mapping: dict[str, str], request) -> dict:
    if not rows:
        raise ValidationError({"rows": ["There are no rows to import."]})
    if len(rows) > MAX_ROWS:
        raise ValidationError({"rows": [f"Import at most {MAX_ROWS} rows at a time."]})
    if not any(field in mapping.values() for field in ("first_name", "last_name", "full_name")):
        raise ValidationError({"mapping": ["Match a column to the prospect's name."]})
    results = _check(rows, mapping, request)
    public = [{k: v for k, v in r.items() if not k.startswith("_")} for r in results]
    return {
        "rows": public,
        "summary": {
            "total": len(results),
            "valid": sum(r["valid"] for r in results),
            "invalid": sum(not r["valid"] for r in results),
            "duplicates": sum(1 for r in results if r["duplicate_of"]),
        },
    }


@transaction.atomic
def import_rows(rows: list, mapping: dict[str, str], request, skip_duplicates: bool = True) -> dict:
    from .serializers import ProspectWriteSerializer
    from .services import create_prospect

    results = _check(rows, mapping, request)
    user = request.user
    lookups = _Lookups(user)
    created = skipped_duplicates = 0
    invalid = sum(not r["valid"] for r in results)
    for r in results:
        if not r["valid"] or (skip_duplicates and r["duplicate_of"]):
            skipped_duplicates += bool(r["valid"] and r["duplicate_of"])
            continue
        data = dict(r["_data"])
        if r["_new_source"]:
            key = r["_new_source"].lower()
            source = lookups.sources.get(key)
            if source is None:
                source = ProspectSource.objects.create(owner=user, name=r["_new_source"], sort_order=len(lookups.sources) + 1)
                lookups.sources[key] = source
            data["source"] = source.pk
        serializer = ProspectWriteSerializer(data=data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        validated = dict(serializer.validated_data)
        interests = validated.pop("insurance_interests", None)
        validated.pop("initial_follow_up", None)
        create_prospect(user=user, data=validated, interests=interests)
        created += 1
    record(user=user, action=AuditAction.PROSPECTS_IMPORTED, entity=None,
           summary=f"Imported {created} prospects from a file",
           changes={"created": created, "skipped_duplicates": skipped_duplicates, "invalid": invalid})
    return {"created": created, "skipped_duplicates": skipped_duplicates, "invalid": invalid}
